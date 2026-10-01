import html
import textwrap
from io import BytesIO
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from fpdf import FPDF


def sanitize_text(text: str) -> str:
    replacements = {"“":"\"", "”":"\"", "‘":"'", "’":"'", "–":"-", "—":"-", "•":"-", "…":"..."}
    for a,b in replacements.items(): text = text.replace(a,b)
    return "".join(ch for ch in text if ch == "\n" or ch == "\t" or ord(ch) >= 32)


def format_html_preview(text: str) -> str:
    parts = []
    for line in html.escape(text).splitlines():
        if not line.strip():
            parts.append("<div style='height:10px'></div>")
        elif line.upper() == line and len(line.strip()) < 80:
            parts.append(f"<h3 style='margin:18px 0 7px'>{line}</h3>")
        elif line.startswith("- "):
            parts.append(f"<div style='margin:5px 0 5px 20px'>• {line[2:]}</div>")
        else:
            parts.append(f"<p style='margin:7px 0'>{line}</p>")
    return "".join(parts)


def format_docx(text: str, doc_type: str) -> bytes:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(.75); sec.bottom_margin = Inches(.75)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("LegalEase"); r.bold = True; r.font.size = Pt(18); r.font.name = "Times New Roman"
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(doc_type.upper()); r.bold = True; r.font.size = Pt(14); r.font.name = "Times New Roman"
    for line in text.splitlines():
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6)
        r = p.add_run(line); r.font.name = "Times New Roman"; r.font.size = Pt(11)
    footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("LegalEase - AI-Powered Legal Document Generator").font.name = "Times New Roman"
    out = BytesIO(); doc.save(out); return out.getvalue()


class LegalEasePDF(FPDF):
    def header(self):
        self.set_font("Times", "B", 14); self.cell(0, 8, "LegalEase", ln=1, align="C"); self.ln(2)
    def footer(self):
        self.set_y(-15); self.set_font("Times", "I", 8); self.cell(0, 10, "LegalEase - AI-Powered Legal Document Generator", align="C")


def format_pdf(text: str, doc_type: str) -> bytes:
    pdf = LegalEasePDF(); pdf.set_auto_page_break(True, 18); pdf.add_page();
    pdf.set_font("Times", "B", 14); pdf.multi_cell(0, 8, doc_type.upper(), align="C"); pdf.ln(4)
    pdf.set_font("Times", "", 11)
    for raw in text.splitlines():
        line = raw.replace("## ", "").replace("**", "")
        if line.strip():
            for chunk in textwrap.wrap(line, width=88, break_long_words=True, break_on_hyphens=False) or [""]:
                pdf.multi_cell(170, 6, chunk)
        else:
            pdf.ln(3)
    return bytes(pdf.output(dest="S"))
